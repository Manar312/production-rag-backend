
from pathlib import Path

from docx import Document

from src.services.document_processors.base import ExtractedContent


class DOCXProcessor:
    def extract(self, file_path: str | Path) -> list[ExtractedContent]:
        path = Path(file_path)
        document = Document(path)
        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f"[Table {table_index}]")
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                parts.append(" | ".join(cells))

        combined_text = "\n".join(parts).strip()

        if not combined_text:
            return []

        return [
            ExtractedContent(
                text=combined_text,
                source=path.name,
            )
        ]
