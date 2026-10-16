
from pathlib import Path

from pptx import Presentation

from src.services.document_processors.base import ExtractedContent


class PPTXProcessor:
    def extract(self, file_path: str | Path) -> list[ExtractedContent]:
        path = Path(file_path)
        presentation = Presentation(path)
        results: list[ExtractedContent] = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        parts.append(text)

                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip().replace("\n", " ")
                            for cell in row.cells
                        ]
                        parts.append(" | ".join(cells))

            combined_text = "\n".join(parts).strip()

            if combined_text:
                results.append(
                    ExtractedContent(
                        text=combined_text,
                        source=path.name,
                        slide_number=slide_number,
                    )
                )

        return results
