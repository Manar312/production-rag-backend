
import csv
from pathlib import Path

from openpyxl import load_workbook

from src.services.document_processors.base import ExtractedContent


class SpreadsheetProcessor:
    def extract(self, file_path: str | Path) -> list[ExtractedContent]:
        path = Path(file_path)
        extension = path.suffix.lower()

        if extension == ".csv":
            return self._extract_csv(path)

        if extension == ".xlsx":
            return self._extract_xlsx(path)

        raise ValueError("Unsupported spreadsheet format.")

    def _extract_csv(self, path: Path) -> list[ExtractedContent]:
        rows: list[str] = []

        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.reader(file)
            for row in reader:
                values = [value.strip() for value in row]
                if any(values):
                    rows.append(" | ".join(values))

        return [
            ExtractedContent(text="\n".join(rows), source=path.name)
        ] if rows else []

    def _extract_xlsx(self, path: Path) -> list[ExtractedContent]:
        results: list[ExtractedContent] = []

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip() if value is not None else ""
                        for value in row
                    ]
                    if any(values):
                        rows.append(" | ".join(values))

                if rows:
                    results.append(
                        ExtractedContent(
                            text="\n".join(rows),
                            source=path.name,
                            sheet_name=sheet.title,
                        )
                    )
        finally:
            workbook.close()

        return results
